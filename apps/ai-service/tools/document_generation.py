from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

FORMAT_META: dict[str, tuple[str, str]] = {
    "markdown": ("text/markdown", ".md"),
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".docx",
    ),
    "pdf": ("application/pdf", ".pdf"),
    "pptx": (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".pptx",
    ),
}


def _slug(value: str) -> str:
    cleaned = re.sub(
        r"[^\w\u4e00-\u9fff-]+", "-", str(value or "").strip(), flags=re.UNICODE
    )
    cleaned = re.sub(r"-+", "-", cleaned).strip("-_")
    return (cleaned or "shinkou-report")[:100]


def _plain_inline(value: str) -> str:
    value = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", value)
    value = re.sub(r"(`+|\*\*|__|\*|_)", "", value)
    return value.strip()


def _markdown_lines(markdown: str) -> list[str]:
    return [
        line.rstrip() for line in str(markdown or "").replace("\r\n", "\n").split("\n")
    ]


def _normalise_markdown(markdown: str, sources: list[dict[str, Any]]) -> str:
    content = str(markdown or "").strip()
    wrapped = re.fullmatch(
        r"\s*```(?:markdown|md)?\s*\n([\s\S]*?)\n```\s*",
        content,
        flags=re.IGNORECASE,
    )
    if wrapped:
        content = wrapped.group(1).strip()
    fenced_blocks: list[str] = []

    def protect(match: re.Match[str]) -> str:
        fenced_blocks.append(match.group(0))
        return f"\x00SHINKOU_CODE_{len(fenced_blocks) - 1}\x00"

    protected = re.sub(r"(```[\s\S]*?```|~~~[\s\S]*?~~~)", protect, content)
    protected = re.sub(r"^(#{1,6})(?=[^#\s])", r"\1 ", protected, flags=re.MULTILINE)
    protected = re.sub(r"^(\s*)([-+])(?=\S)", r"\1\2 ", protected, flags=re.MULTILINE)
    protected = re.sub(r"^(\s*)(\d+[.)])(?=\S)", r"\1\2 ", protected, flags=re.MULTILINE)

    def restore(match: re.Match[str]) -> str:
        return fenced_blocks[int(match.group(1))]

    content = re.sub(r"\x00SHINKOU_CODE_(\d+)\x00", restore, protected)
    if not sources:
        return content + ("\n" if content else "")
    source_lines = ["", "## 来源", ""]
    for item in sources[:30]:
        title = str(
            item.get("title") or item.get("name") or item.get("source") or "未命名来源"
        ).strip()
        source_id = str(item.get("id") or "").strip()
        url = str(item.get("url") or "").strip()
        label = f"[{title}]({url})" if url else title
        suffix = f"（{source_id}）" if source_id else ""
        source_lines.append(f"- {label}{suffix}")
    return (content + "\n" if content else "") + "\n".join(source_lines) + "\n"


def _set_run_font(
    run: Any,
    name: str = "Microsoft YaHei",
    size: int = 11,
    bold: bool = False,
    italic: bool = False,
    underline: bool = False,
    strike: bool = False,
    color: str | None = None,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.strike = strike
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_hyperlink(paragraph: Any, label: str, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    part = paragraph.part
    relationship = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Microsoft YaHei")
    r_fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    properties.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "147d78")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_markdown_runs(paragraph: Any, text: str) -> None:
    pattern = re.compile(
        r"(?P<image>!\[(?P<image_label>[^\]]*)\]\((?P<image_url>[^)\s]+)(?:\s+\"[^\"]*\")?\))|"
        r"(?P<link>\[(?P<link_label>[^\]]+)\]\((?P<link_url>https?://[^)]+)\))|"
        r"(?P<strong>\*\*(?P<strong_text>.+?)\*\*|__(?P<strong_alt>.+?)__)|"
        r"(?P<strike>~~(?P<strike_text>.+?)~~)|"
        r"(?P<code>`(?P<code_text>[^`]+)`)|"
        r"(?P<em>\*(?P<em_text>[^*]+?)\*|(?<!\w)_(?P<em_alt>[^_]+?)_(?!\w))"
    )
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            _set_run_font(run)
        if match.group("link_label") and match.group("link_url"):
            _add_hyperlink(paragraph, match.group("link_label"), match.group("link_url"))
        elif match.group("image_label"):
            run = paragraph.add_run(match.group("image_label"))
            _set_run_font(run, italic=True, color="666666")
        elif match.group("strong_text") or match.group("strong_alt"):
            run = paragraph.add_run(match.group("strong_text") or match.group("strong_alt"))
            _set_run_font(run, bold=True)
        elif match.group("strike_text"):
            run = paragraph.add_run(match.group("strike_text"))
            _set_run_font(run, strike=True)
        elif match.group("code_text"):
            run = paragraph.add_run(match.group("code_text"))
            _set_run_font(run, name="Consolas", size=9, color="7A3E00")
        elif match.group("em_text") or match.group("em_alt"):
            run = paragraph.add_run(match.group("em_text") or match.group("em_alt"))
            _set_run_font(run, italic=True)
        else:
            run = paragraph.add_run(match.group(0))
            _set_run_font(run)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        _set_run_font(run)


def _split_table_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|") and not value.endswith("\\|"):
        value = value[:-1]
    cells: list[str] = []
    current: list[str] = []
    escaped = False
    for character in value:
        if character == "|" and not escaped:
            cells.append("".join(current).replace("\\|", "|").strip())
            current = []
            continue
        current.append(character)
        escaped = character == "\\" and not escaped
        if character != "\\":
            escaped = False
    cells.append("".join(current).replace("\\|", "|").strip())
    return cells


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return len(cells) >= 1 and all(
        bool(re.fullmatch(r":?-{3,}:?", cell.replace(" ", ""))) for cell in cells
    )


def _is_block_start(lines: list[str], index: int) -> bool:
    value = lines[index]
    if re.match(r"^\s*(?:`{3,}|~{3,})", value):
        return True
    if re.match(r"^\s*#{1,6}\s+\S", value):
        return True
    if re.match(r"^\s*>\s?", value):
        return True
    if re.match(r"^\s*(?:[-+*]|\d+[.)])\s+\S", value):
        return True
    if re.match(r"^\s*(?:[-*_]\s*){3,}$", value):
        return True
    return index + 1 < len(lines) and "|" in value and _is_table_separator(lines[index + 1])


def _parse_markdown_blocks(markdown: str) -> list[dict[str, Any]]:
    """Parse Markdown blocks before writing native DOCX nodes."""

    lines = _markdown_lines(markdown)
    blocks: list[dict[str, Any]] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        fence = re.match(r"^\s*(`{3,}|~{3,})\s*([\w+-]*)\s*$", line)
        if fence:
            marker = fence.group(1)[0]
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not re.match(
                rf"^\s*{re.escape(marker)}{{{len(fence.group(1))},}}\s*$", lines[index]
            ):
                code_lines.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            blocks.append({"kind": "code", "language": fence.group(2), "text": "\n".join(code_lines)})
            continue

        heading = re.match(r"^\s*(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            blocks.append({"kind": "heading", "level": len(heading.group(1)), "text": heading.group(2).strip()})
            index += 1
            continue

        if index + 1 < len(lines) and "|" in line and _is_table_separator(lines[index + 1]):
            headers = _split_table_row(line)
            alignments = _split_table_row(lines[index + 1])
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                rows.append(_split_table_row(lines[index]))
                index += 1
            blocks.append({"kind": "table", "headers": headers, "alignments": alignments, "rows": rows})
            continue

        if re.match(r"^\s*(?:[-*_]\s*){3,}$", line):
            blocks.append({"kind": "rule"})
            index += 1
            continue

        if re.match(r"^\s*>\s?", line):
            quote_lines: list[str] = []
            while index < len(lines):
                match = re.match(r"^\s*>\s?(.*)$", lines[index])
                if not match:
                    break
                quote_lines.append(match.group(1))
                index += 1
            blocks.append({"kind": "quote", "text": "\n".join(quote_lines)})
            continue

        list_match = re.match(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$", line)
        if list_match:
            base_indent = len(list_match.group(1).replace("\t", "    "))
            items: list[dict[str, Any]] = []
            while index < len(lines):
                current = lines[index]
                match = re.match(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$", current)
                if match:
                    indent = len(match.group(1).replace("\t", "    "))
                    if indent < base_indent:
                        break
                    items.append(
                        {
                            "ordered": match.group(2)[0].isdigit(),
                            "level": max(0, (indent - base_indent) // 2),
                            "text": match.group(3).strip(),
                        }
                    )
                    index += 1
                    continue
                if items and current.strip() and len(current) - len(current.lstrip()) > base_indent:
                    items[-1]["text"] += " " + current.strip()
                    index += 1
                    continue
                break
            blocks.append({"kind": "list", "items": items})
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and lines[index].strip():
            if _is_block_start(lines, index):
                break
            paragraph_lines.append(lines[index].strip())
            index += 1
        blocks.append({"kind": "paragraph", "text": " ".join(paragraph_lines)})
    return blocks


def _add_soft_break_text(paragraph: Any, text: str) -> None:
    for index, line in enumerate(str(text).split("\n")):
        if index:
            paragraph.add_run().add_break()
        _add_markdown_runs(paragraph, line)


def _set_paragraph_shading(paragraph: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _write_docx(path: Path, markdown: str) -> None:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color in (
        ("Title", 22, "174b4a"),
        ("Heading 1", 16, "147d78"),
        ("Heading 2", 13, "286a68"),
        ("Heading 3", 11, "286a68"),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    if "Quote" not in styles:
        styles.add_style("Quote", WD_STYLE_TYPE.PARAGRAPH)
    quote = styles["Quote"]
    quote.font.name = "Microsoft YaHei"
    quote.font.size = Pt(10.5)
    quote.font.italic = True
    quote.font.color.rgb = RGBColor.from_string("52656A")
    quote._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    quote.paragraph_format.left_indent = Inches(0.3)
    quote.paragraph_format.space_after = Pt(6)
    if "Code Block" not in styles:
        styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = RGBColor.from_string("4A3424")
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code_style.paragraph_format.left_indent = Inches(0.25)
    code_style.paragraph_format.space_after = Pt(6)

    first_heading = True
    for block in _parse_markdown_blocks(markdown):
        kind = block["kind"]
        if kind == "heading":
            level = min(6, int(block["level"]))
            style_name = "Title" if level == 1 and first_heading else f"Heading {min(level, 3)}"
            paragraph = document.add_paragraph(style=style_name)
            _add_markdown_runs(paragraph, block["text"])
            first_heading = False
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            _add_soft_break_text(paragraph, block["text"])
        elif kind == "quote":
            paragraph = document.add_paragraph(style="Quote")
            _add_soft_break_text(paragraph, block["text"])
            _set_paragraph_shading(paragraph, "EAF3F2")
        elif kind == "code":
            paragraph = document.add_paragraph(style="Code Block")
            code_lines = str(block["text"]).split("\n") or [""]
            for index, code_line in enumerate(code_lines):
                if index:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(code_line)
                _set_run_font(run, name="Consolas", size=9, color="4A3424")
            _set_paragraph_shading(paragraph, "F5EEE8")
        elif kind == "list":
            for item in block["items"]:
                level = min(3, int(item["level"]))
                style_name = "List Number" if item["ordered"] else "List Bullet"
                if level:
                    candidate = f"{style_name} {level + 1}"
                    if candidate in styles:
                        style_name = candidate
                paragraph = document.add_paragraph(style=style_name)
                _add_markdown_runs(paragraph, item["text"])
        elif kind == "table":
            headers = list(block["headers"])
            width = max(1, len(headers), *(len(row) for row in block["rows"]))
            table = document.add_table(rows=1, cols=width)
            try:
                table.style = "Light Shading Accent 1"
            except KeyError:
                table.style = "Table Grid"
            for column in range(width):
                cell = table.rows[0].cells[column]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                _add_markdown_runs(paragraph, headers[column] if column < len(headers) else "")
                for run in paragraph.runs:
                    run.bold = True
            for row in block["rows"]:
                cells = table.add_row().cells
                for column in range(width):
                    cells[column].text = ""
                    _add_markdown_runs(cells[column].paragraphs[0], row[column] if column < len(row) else "")
        elif kind == "rule":
            paragraph = document.add_paragraph()
            properties = paragraph._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "B8D6D2")
            border.append(bottom)
            properties.append(border)
    document.save(path)


def _add_slide_text(
    slide: Any,
    title: str,
    body: list[str],
    *,
    title_size: int = 30,
    body_size: int = 18,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    title_box = slide.shapes.add_textbox(
        Inches(0.7), Inches(0.45), Inches(12), Inches(0.75)
    )
    title_frame = title_box.text_frame
    title_frame.clear()
    title_frame.word_wrap = True
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.alignment = PP_ALIGN.LEFT
    run = title_paragraph.add_run()
    run.text = title
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(title_size)
    run.font.bold = True
    run.font.color.rgb = RGBColor(24, 91, 88)
    body_box = slide.shapes.add_textbox(
        Inches(0.85), Inches(1.45), Inches(11.7), Inches(5.45)
    )
    frame = body_box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, item in enumerate(body[:12]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = _plain_inline(item)
        paragraph.level = 0
        paragraph.space_after = Pt(10)
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(body_size)
            run.font.color.rgb = RGBColor(46, 59, 67)


def _write_pptx(
    path: Path, title: str, markdown: str, sources: list[dict[str, Any]]
) -> None:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    cover = presentation.slides.add_slide(blank)
    _add_slide_text(cover, title, ["研究整理与行动摘要"], title_size=36, body_size=22)
    sections: list[tuple[str, list[str]]] = []
    current_title = "摘要"
    current_body: list[str] = []
    for raw in _markdown_lines(markdown):
        line = raw.strip()
        heading = re.match(r"^#{1,3}\s+(.+)$", line)
        if heading:
            if current_body:
                sections.append((current_title, current_body))
            current_title = _plain_inline(heading.group(1))
            current_body = []
        elif line:
            current_body.append(line)
    if current_body:
        sections.append((current_title, current_body))
    for section_title, body in sections[:12]:
        slide = presentation.slides.add_slide(blank)
        _add_slide_text(slide, section_title, body)
    if sources:
        slide = presentation.slides.add_slide(blank)
        source_lines = []
        for item in sources[:20]:
            label = str(
                item.get("title")
                or item.get("name")
                or item.get("source")
                or "未命名来源"
            )
            url = str(item.get("url") or "").strip()
            source_lines.append(f"• {label}{f' — {url}' if url else ''}")
        _add_slide_text(slide, "来源", source_lines, title_size=26, body_size=14)
    presentation.save(path)


def _convert_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    command = shutil.which("soffice") or shutil.which("libreoffice")
    if not command:
        raise RuntimeError("未找到 LibreOffice/soffice，无法把 DOCX 转换为 PDF")
    profile_dir = output_dir / ".libreoffice-profile"
    profile_dir.mkdir(parents=True, exist_ok=True)
    profile_uri = profile_dir.as_uri()
    completed = subprocess.run(
        [
            command,
            "--headless",
            "--norestore",
            "--nodefault",
            "--nolockcheck",
            "--nofirststartwizard",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx_path),
        ],
        capture_output=True,
        text=True,
        timeout=90,
        check=False,
    )
    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    if completed.returncode != 0 or not pdf_path.exists():
        detail = (completed.stderr or completed.stdout or "转换没有产生 PDF").strip()
        raise RuntimeError(detail[-500:])
    return pdf_path


class DocumentGenerationTool:
    """Render the final Markdown into native, format-specific document structures.

    Markdown is only the source representation. It must never be copied as the
    visible body of a DOCX/PDF/PPTX artifact.
    """

    def generate(self, **input_data: Any) -> dict[str, Any]:
        title = str(input_data.get("title") or "Shinkou 研究报告").strip()[:200]
        markdown = str(input_data.get("markdown") or "").strip()
        formats = [str(item).lower().strip() for item in input_data.get("formats", [])]
        formats = list(dict.fromkeys(item for item in formats if item in FORMAT_META))
        sources = (
            input_data.get("sources")
            if isinstance(input_data.get("sources"), list)
            else []
        )
        output_dir = (
            Path(str(input_data.get("output_dir") or "")).expanduser().resolve()
        )
        if not markdown:
            raise ValueError("markdown 不能为空")
        if not formats:
            raise ValueError("formats 至少需要包含一种输出格式")
        output_dir.mkdir(parents=True, exist_ok=True)
        document_markdown = _normalise_markdown(markdown, sources)
        base_name = _slug(title)
        artifacts: list[dict[str, Any]] = []
        warnings: list[str] = []
        docx_path: Path | None = None

        if "markdown" in formats:
            path = output_dir / f"{base_name}.md"
            path.write_text(document_markdown, encoding="utf-8")
            artifacts.append(self._artifact(path, "markdown"))
        if "docx" in formats or "pdf" in formats:
            docx_path = output_dir / f"{base_name}.docx"
            _write_docx(docx_path, document_markdown)
            if "docx" in formats:
                artifacts.append(self._artifact(docx_path, "docx"))
        if "pdf" in formats and docx_path is not None:
            try:
                artifacts.append(
                    self._artifact(_convert_to_pdf(docx_path, output_dir), "pdf")
                )
            except Exception as exc:
                warnings.append(f"PDF 生成失败：{str(exc)[:300]}")
        if "pptx" in formats:
            path = output_dir / f"{base_name}.pptx"
            _write_pptx(path, title, markdown, sources)
            artifacts.append(self._artifact(path, "pptx"))
        return {"artifacts": artifacts, "warnings": warnings}

    @staticmethod
    def _artifact(path: Path, format_name: str) -> dict[str, Any]:
        mime_type, _ = FORMAT_META[format_name]
        return {
            "fileName": path.name,
            "mimeType": mime_type,
            "fileSize": path.stat().st_size,
            "path": str(path),
            "format": format_name,
        }
