from pathlib import Path

from documents.parser import DocumentParser
from tools.document_generation import DocumentGenerationTool


def test_document_generation_bundle_writes_markdown_docx_and_pptx(tmp_path: Path) -> None:
    result = DocumentGenerationTool().generate(
        title="图形学进展",
        markdown="# 图形学进展\n\n## 重点\n\n- 实时渲染\n- 数字人",
        formats=["markdown", "docx", "pptx"],
        output_dir=str(tmp_path),
        sources=[{"id": "W1", "title": "官方来源", "url": "https://example.com"}],
    )

    names = {item["fileName"] for item in result["artifacts"]}
    assert names == {"图形学进展.md", "图形学进展.docx", "图形学进展.pptx"}
    assert "[官方来源](https://example.com)" in (tmp_path / "图形学进展.md").read_text(encoding="utf-8")
    assert all((tmp_path / name).stat().st_size > 0 for name in names)
    assert result["warnings"] == []


def test_docx_generation_renders_markdown_as_native_structure(tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    result = DocumentGenerationTool().generate(
        title="格式化报告",
        markdown=(
            "# 主标题\n\n"
            "## 结论\n\n"
            "这是 **重点**、*强调*、`代码` 和 [来源](https://example.com)。\n\n"
            "- 第一项\n- 第二项\n\n"
            "| 指标 | 结果 |\n| --- | --- |\n| 覆盖率 | 95% |\n\n"
            "> 这是引用\n\n"
            "```python\nprint('ok')\n```"
        ),
        formats=["docx"],
        output_dir=str(tmp_path),
    )

    assert result["warnings"] == []
    document = DocxDocument(tmp_path / "格式化报告.docx")
    paragraph_texts = [paragraph.text for paragraph in document.paragraphs]
    assert paragraph_texts[0] == "主标题"
    assert "重点" in paragraph_texts[2]
    assert "第一项" in paragraph_texts[3]
    assert all(not text.lstrip().startswith(("# ", "- ", "> ", "```")) for text in paragraph_texts)
    assert document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[1].style.name.startswith("Heading")
    assert any(run.bold for run in document.paragraphs[2].runs)
    assert any(run.italic for run in document.paragraphs[2].runs)
    assert len(document.tables) == 1

    parsed = DocumentParser().parse_bytes(
        (tmp_path / "格式化报告.docx").read_bytes(),
        file_name="格式化报告.docx",
    )
    assert "# 主标题" in parsed.text
    assert "- 第一项" in parsed.text
    assert "[来源](https://example.com)" in parsed.text
    assert parsed.metadata["headingCount"] == 2
    assert parsed.metadata["listCount"] == 2
    assert parsed.metadata["tableCount"] == 1
