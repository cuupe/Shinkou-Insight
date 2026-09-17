import pytest
from io import BytesIO

from documents.chunker import DocumentChunker
from documents.parser import DocumentParser
from embeddings.providers import HashEmbeddingProvider
from graph.store import GraphExtractor, MemoryGraphStore
from rag.indexer import InMemoryKnowledgeStore, KnowledgeIndexer


@pytest.mark.asyncio
async def test_document_parse_chunk_embed_and_graph_pipeline():
    embedding = HashEmbeddingProvider(dimension=64)
    store = InMemoryKnowledgeStore(embedding)
    graph = MemoryGraphStore()
    indexer = KnowledgeIndexer(DocumentParser(), DocumentChunker(chunk_size=40, chunk_overlap=8), embedding, store, graph, GraphExtractor())
    result = await indexer.index_bytes(data="# Architecture\nKafka and PostgreSQL are used by the service.".encode(), file_name="architecture.md", mime_type="text/markdown", workspace_id=7, project_id=8, asset_id=9)
    assert result["chunk_count"] >= 1
    assert result["embedding_dimension"] == 64
    assert result["graph_entities"] >= 2
    items = await store.retrieve(workspace_id=7, project_id=8, question="Kafka PostgreSQL", top_k=3, retrieval_mode="HYBRID")
    assert items and items[0].asset_id == 9
    nodes, edges = await graph.search(workspace_id=7, project_id=8, query="Kafka", limit=5)
    assert nodes and any(node.name == "Kafka" for node in nodes)


@pytest.mark.asyncio
async def test_file_parser_rejects_unsupported_type():
    with pytest.raises(ValueError, match="Unsupported document type"):
        DocumentParser().parse_bytes(b"binary", file_name="archive.zip")


def test_chunker_keeps_markdown_section_title_with_each_section():
    document = DocumentParser().parse_bytes(
        b"# Overview\nFirst paragraph.\n\n## Details\nSecond paragraph.",
        file_name="notes.md",
        mime_type="text/markdown",
    ).documents[0]

    chunks = DocumentChunker(chunk_size=200, chunk_overlap=20).split([document])

    assert [chunk.section_title for chunk in chunks] == ["Overview", "Details"]
    assert all("paragraph" in chunk.content for chunk in chunks)


def test_docx_parser_extracts_tables_and_keeps_document_metadata():
    from docx import Document as DocxDocument

    source = DocxDocument()
    source.add_heading("项目结论", level=1)
    source.add_paragraph("正文内容")
    table = source.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "结果"
    buffer = BytesIO()
    source.save(buffer)

    parsed = DocumentParser().parse_bytes(buffer.getvalue(), file_name="report.docx")

    assert parsed.documents
    assert "项目结论" in parsed.text
    assert "指标\t结果" in parsed.text
    assert parsed.metadata["tableCount"] == 1


def test_media_parser_degrades_with_explicit_tool_warning():
    parsed = DocumentParser().parse_bytes(b"not-an-image", file_name="diagram.png", mime_type="image/png")

    assert parsed.metadata["kind"] == "image"
    assert parsed.documents
    assert parsed.warnings or "未识别到" in parsed.text


@pytest.mark.asyncio
async def test_indexer_applies_project_chunking_config():
    embedding = HashEmbeddingProvider(dimension=32)
    store = InMemoryKnowledgeStore(embedding)
    indexer = KnowledgeIndexer(DocumentParser(), DocumentChunker(), embedding, store)
    result = await indexer.index_bytes(
        data=("第一段内容。" * 30).encode(),
        file_name="config.txt",
        mime_type="text/plain",
        workspace_id=1,
        project_id=2,
        asset_id=3,
        chunking={"strategy": "fixed", "chunk_size": 40, "chunk_overlap": 0, "preserve_sections": False},
    )

    assert result["chunk_count"] > 1
    assert all(len(item.chunk.content) <= 40 for item in store.items)
